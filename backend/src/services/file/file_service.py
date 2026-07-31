import asyncio
import uuid
from pathlib import Path

from fastapi import UploadFile
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.config import settings
from src.utils import file_parser
from src.utils.user_dir import resolve_username as _resolve_username

import logging

logger = logging.getLogger(__name__)

UPLOAD_DIR = Path(settings.upload_dir)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def get_user_upload_dir(user_id: int | str) -> Path:
    """Return the per-user upload directory, creating it if needed."""
    user_dir = UPLOAD_DIR / _resolve_username(user_id)
    user_dir.mkdir(parents=True, exist_ok=True)
    return user_dir

# Max 100MB per file
MAX_FILE_SIZE = 100 * 1024 * 1024

DOCUMENT_TYPES = {
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    ".pdf": {"application/pdf"},
}

IMAGE_TYPES = {
    ".png": {"image/png"},
    ".jpg": {"image/jpeg"},
    ".jpeg": {"image/jpeg"},
    ".webp": {"image/webp"},
    ".svg": {"image/svg+xml"},
}

SUPPORTED_TYPES = DOCUMENT_TYPES | IMAGE_TYPES


def _get_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def _is_supported(filename: str, *, allow_images: bool = True) -> bool:
    ext = _get_extension(filename)
    supported = SUPPORTED_TYPES if allow_images else DOCUMENT_TYPES
    return ext in supported


def _validate_file(filename: str, size: int | None, content_type: str | None, *, allow_images: bool = True) -> str | None:
    """Return error message if invalid, None if OK."""
    if size is not None and size > MAX_FILE_SIZE:
        return "File exceeds 100MB limit"
    if not _is_supported(filename, allow_images=allow_images):
        ext = _get_extension(filename)
        return f"Unsupported file type: {ext}"
    return None


async def save_file(file: UploadFile, *, allow_images: bool = True, user_id: int | str = 1) -> tuple[str, str]:
    """Save uploaded file and return (stored_filename, original_name)."""
    error = _validate_file(file.filename or "file", file.size, file.content_type, allow_images=allow_images)
    if error:
        raise ValueError(error)

    # Read content first before opening the output file to avoid empty writes
    content = await file.read()
    if not content:
        raise ValueError("Uploaded file is empty")

    ext = _get_extension(file.filename or "file")
    stored_name = f"{uuid.uuid4().hex}{ext}"
    user_dir = get_user_upload_dir(user_id)
    path = user_dir / stored_name

    path.write_bytes(content)

    return stored_name, file.filename or "file"


def delete_uploaded_file(stored_name: str, user_id: int | str) -> bool:
    """Delete a single uploaded file by stored name and user directory. Returns True if deleted."""
    path = get_user_upload_dir(user_id) / stored_name
    try:
        if path.exists() and path.is_file():
            path.unlink()
            logger.info("Deleted uploaded file: %s", path)
            return True
    except Exception:
        logger.exception("Failed to delete uploaded file: %s", path)
    return False


async def is_hidden_soft_deleted_file(
    db: AsyncSession,
    *,
    filename: str,
    user_id: int,
    username: str,
) -> bool:
    """仅当文件只被软删除的文件库记录引用时隐藏。"""
    from src.database.models import BlogPost, Conversation, FileDocument, Message

    states = (
        await db.execute(
            select(FileDocument.deleted_at).where(
                FileDocument.user_id == str(user_id),
                FileDocument.file_path == filename,
            )
        )
    ).scalars().all()
    if not states or any(deleted_at is None for deleted_at in states):
        return False

    references = {
        filename,
        f"/api/v1/uploads/{filename}",
        f"/api/v1/public/uploads/{username}/{filename}",
        f"/api/v1/blog/cover/{filename}",
    }
    message_ref = await db.execute(
        select(Message.id)
        .join(Conversation, Conversation.id == Message.conversation_id)
        .where(
            Conversation.user_id == user_id,
            Conversation.deleted_at.is_(None),
            or_(Message.image_url.in_(references), Message.file_url.in_(references)),
        )
        .limit(1)
    )
    if message_ref.first() is not None:
        return False

    blog_ref = await db.execute(
        select(BlogPost.id).where(
            BlogPost.user_id == user_id,
            BlogPost.deleted_at.is_(None),
            BlogPost.cover_image.in_(references),
        ).limit(1)
    )
    return blog_ref.first() is None


async def vectorize_and_store(
    stored_filename: str,
    collection_name: str,
    original_name: str | None = None,
    category_id: int | None = None,
    user_id: int | str = "default_user",
    progress_reporter=None,
) -> list[str]:
    """Parse, chunk, embed, and store documents in the vector store.

    Returns:
        List of chunk contents that were stored.
    """
    from src.services.rag.embedding_service import get_embeddings
    from src.services.rag.vector_store import add_documents

    from src.utils.chunker import chunk_text

    async def report(stage: str, completed: int, total: int, unit: str) -> None:
        if progress_reporter:
            await progress_reporter(stage, completed, total, unit)

    logger.info(
        "KB vectorization parse started: stored_name=%s original_name=%s user_id=%s category_id=%s",
        stored_filename,
        original_name,
        user_id,
        category_id,
    )
    loop = asyncio.get_running_loop()

    def parser_progress(completed: int, total: int, unit: str) -> None:
        if progress_reporter:
            asyncio.run_coroutine_threadsafe(report("parse", completed, total, unit), loop).result()

    path = get_user_upload_dir(user_id) / stored_filename
    await report("parse", 0, 1, "operation")
    text = await asyncio.to_thread(file_parser.parse_path, path, parser_progress)
    logger.info(
        "KB vectorization parse completed: stored_name=%s text_length=%s",
        stored_filename,
        len(text),
    )

    await report("chunk", 0, 1, "operation")
    chunks = await asyncio.to_thread(chunk_text, text)
    await report("chunk", 1, 1, "operation")
    logger.info(
        "KB vectorization chunking completed: stored_name=%s chunks=%s",
        stored_filename,
        len(chunks),
    )
    if not chunks:
        logger.warning("KB vectorization produced no chunks: stored_name=%s", stored_filename)
        raise ValueError("Document contains no indexable text")

    logger.info("KB vectorization embedding started: stored_name=%s chunks=%s", stored_filename, len(chunks))

    async def embedding_progress(completed: int, total: int, unit: str) -> None:
        await report("embedding", completed, total, unit)

    await report("embedding", 0, len(chunks), "chunk")
    embeddings = await get_embeddings(chunks, progress_callback=embedding_progress)
    if len(embeddings) != len(chunks):
        raise ValueError(f"Embedding count mismatch: chunks={len(chunks)} embeddings={len(embeddings)}")
    logger.info("KB vectorization embedding completed: stored_name=%s embeddings=%s", stored_filename, len(embeddings))

    file_name = original_name or stored_filename
    file_type = _get_extension(file_name).lstrip(".")
    metadata_list = []
    await report("metadata", 0, len(chunks), "chunk")
    for index in range(len(chunks)):
        metadata = {
            "source": file_name,
            "file_name": file_name,
            "original_name": file_name,
            "stored_name": stored_filename,
            "file_type": file_type,
            "collection_name": collection_name,
            "user_id": user_id,
            "chunk_index": index,
            "total_chunks": len(chunks),
            "chunk_id": f"{stored_filename}:{index}",
        }
        if category_id is not None:
            metadata["category_id"] = category_id
        metadata_list.append(metadata)
        await report("metadata", index + 1, len(chunks), "chunk")

    logger.info(
        "KB vectorization chroma upsert started: stored_name=%s collection=%s chunks=%s",
        stored_filename,
        collection_name,
        len(chunks),
    )
    async def vector_progress(completed: int, total: int, unit: str) -> None:
        await report("vector_store", completed, total, unit)

    await report("vector_store", 0, len(chunks), "chunk")
    await add_documents(
        collection_name,
        chunks,
        metadata_list,
        embeddings=embeddings,
        progress_callback=vector_progress,
    )
    logger.info(
        "KB vectorization chroma upsert completed: stored_name=%s collection=%s chunks=%s",
        stored_filename,
        collection_name,
        len(chunks),
    )

    return chunks


async def vectorize_text_and_store(
    text: str,
    collection_name: str,
    *,
    source_id: str,
    original_name: str | None = None,
    user_id: int | str = "default_user",
    resource_type: str = "blog_post",
    category_id: int | None = None,
    progress_reporter=None,
) -> list[str]:
    """对纯文本（如博客 Markdown 正文）分块、向量化、写入向量库。

    与 vectorize_and_store 的差异：跳过文件解析（直接用传入 text），metadata 的
    stored_name/chunk_id 用 source_id（资源稳定标识，如 "blog_post:123"），并附带
    resource_type 便于检索侧按资源类型过滤。
    """
    from src.services.rag.embedding_service import get_embeddings
    from src.services.rag.vector_store import add_documents
    from src.utils.chunker import chunk_text

    async def report(stage: str, completed: int, total: int, unit: str) -> None:
        if progress_reporter:
            await progress_reporter(stage, completed, total, unit)

    logger.info(
        "KB text vectorization started: source_id=%s resource_type=%s user_id=%s",
        source_id,
        resource_type,
        user_id,
    )

    # 对齐 index_v1 权重表（含 parse）：文本无需解析，瞬间完成 parse 段。
    await report("parse", 1, 1, "operation")

    await report("chunk", 0, 1, "operation")
    chunks = await asyncio.to_thread(chunk_text, text)
    await report("chunk", 1, 1, "operation")
    logger.info(
        "KB text vectorization chunking completed: source_id=%s chunks=%s",
        source_id,
        len(chunks),
    )
    if not chunks:
        logger.warning("KB text vectorization produced no chunks: source_id=%s", source_id)
        raise ValueError("Document contains no indexable text")

    async def embedding_progress(completed: int, total: int, unit: str) -> None:
        await report("embedding", completed, total, unit)

    await report("embedding", 0, len(chunks), "chunk")
    embeddings = await get_embeddings(chunks, progress_callback=embedding_progress)
    if len(embeddings) != len(chunks):
        raise ValueError(f"Embedding count mismatch: chunks={len(chunks)} embeddings={len(embeddings)}")

    display_name = original_name or source_id
    metadata_list = []
    await report("metadata", 0, len(chunks), "chunk")
    for index in range(len(chunks)):
        metadata = {
            "source": display_name,
            "original_name": display_name,
            "stored_name": source_id,
            "resource_type": resource_type,
            "collection_name": collection_name,
            "user_id": user_id,
            "chunk_index": index,
            "total_chunks": len(chunks),
            "chunk_id": f"{source_id}:{index}",
        }
        if category_id is not None:
            metadata["category_id"] = category_id
        metadata_list.append(metadata)
        await report("metadata", index + 1, len(chunks), "chunk")

    async def vector_progress(completed: int, total: int, unit: str) -> None:
        await report("vector_store", completed, total, unit)

    await report("vector_store", 0, len(chunks), "chunk")
    await add_documents(
        collection_name,
        chunks,
        metadata_list,
        embeddings=embeddings,
        progress_callback=vector_progress,
    )
    logger.info(
        "KB text vectorization completed: source_id=%s collection=%s chunks=%s",
        source_id,
        collection_name,
        len(chunks),
    )
    return chunks


# ---- File Preview (docx/xlsx → HTML) ----


async def convert_to_html(stored_filename: str, *, user_id: int | str = 1) -> str:
    """Convert a stored file to HTML for browser preview.

    Supports docx → HTML, xlsx → HTML table.
    PDF is not converted (served as-is, browser renders natively).
    """
    path = get_user_upload_dir(user_id) / stored_filename
    ext = _get_extension(stored_filename)
    if ext == ".docx":
        return _docx_to_html(path)
    elif ext == ".xlsx":
        return _xlsx_to_html(path)
    else:
        raise ValueError(f"Preview not supported for: {ext}")


def _docx_to_html(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    html_parts: list[str] = ['<div class="docx-preview">']
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            html_parts.append("<p>&nbsp;</p>")
            continue
        style = para.style.name.lower() if para.style else ""
        if "heading 1" in style:
            html_parts.append(f"<h1>{_escape_html(text)}</h1>")
        elif "heading 2" in style:
            html_parts.append(f"<h2>{_escape_html(text)}</h2>")
        elif "heading 3" in style:
            html_parts.append(f"<h3>{_escape_html(text)}</h3>")
        else:
            html_parts.append(f"<p>{_escape_html(text)}</p>")
    for table in doc.tables:
        html_parts.append('<table border="1" cellpadding="4" style="border-collapse:collapse;width:100%">')
        for i, row in enumerate(table.rows):
            tag = "th" if i == 0 else "td"
            html_parts.append(f"<tr>{''.join(f'<{tag}>{_escape_html(cell.text.strip())}</{tag}>' for cell in row.cells)}</tr>")
        html_parts.append("</table>")
    html_parts.append("</div>")
    return "\n".join(html_parts)


def _xlsx_to_html(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    html_parts: list[str] = ['<div class="xlsx-preview">']
    for ws in wb.worksheets:
        html_parts.append(f"<h3>{_escape_html(ws.title)}</h3>")
        html_parts.append('<table border="1" cellpadding="4" style="border-collapse:collapse;width:100%">')
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            tag = "th" if row_idx == 0 else "td"
            cells = "".join(f"<{tag}>{_escape_html(str(c) if c is not None else '')}</{tag}>" for c in row)
            html_parts.append(f"<tr>{cells}</tr>")
        html_parts.append("</table>")
    html_parts.append("</div>")
    wb.close()
    return "\n".join(html_parts)


def _escape_html(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
