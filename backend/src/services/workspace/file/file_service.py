import asyncio
import time
import uuid
from pathlib import Path, PurePosixPath

from fastapi import UploadFile
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.core.exceptions import ValidationFailedError
from src.core.path_guard import workspace_dir, workspace_path
from src.core.workspace_path import validate_workspace_relative_path
from src.utils import file_parser

import logging

logger = logging.getLogger(__name__)

UPLOAD_ROOT = "uploads"


def normalize_workspace_file_path(stored_path: str) -> str:
    """Return the canonical workspace-relative path for a stored upload."""
    try:
        validated = validate_workspace_relative_path(stored_path)
    except ValidationFailedError as exc:
        raise ValueError("Invalid stored upload path") from exc
    return validated if PurePosixPath(validated).parts[0] == UPLOAD_ROOT else f"{UPLOAD_ROOT}/{validated}"


def get_user_upload_dir(user_id: int) -> Path:
    """Return the per-user upload directory path (does NOT create it).

    纯取路径：读接口（预览/下载/公共图片）不应有建目录副作用，否则未校验的 username
    会被滥用来制造大量空目录。写入场景（save_file 等）自行 mkdir。
    """
    return workspace_dir(user_id) / UPLOAD_ROOT


def get_uploaded_file_path(
    user_id: int,
    stored_path: str,
    *,
    mode: str = "read",
) -> Path:
    """Resolve one document path within the user's real workspace."""
    return workspace_path(
        user_id,
        normalize_workspace_file_path(stored_path),
        mode="write" if mode == "write" else "read",
    )


async def _store_document_chunks(
    *,
    collection_name: str,
    chunks: list[str],
    metadata_list: list[dict],
    embeddings: list[list[float]],
    user_id: int | str,
    progress_callback=None,
) -> None:
    if not isinstance(user_id, int):
        raise ValueError("Document indexing requires a numeric user ID")

    from src.services.memory.graph_store import add_document_chunks
    from src.services.infra.embeddings.embedding_service import get_embedding_collection_suffix

    await add_document_chunks(
        user_id=user_id,
        collection_name=collection_name,
        stored_name=str(metadata_list[0]["stored_name"]),
        chunks=chunks,
        embeddings=embeddings,
        metadata_list=metadata_list,
        embedding_model=get_embedding_collection_suffix(),
        vector_dim=len(embeddings[0]),
        progress_callback=progress_callback,
    )


MAX_FILE_SIZE = 100 * 1024 * 1024

DOCUMENT_TYPES = {
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    ".pdf": {"application/pdf"},
}

IMAGE_TYPES = {
    ".png": {"image/png"},
    ".jpg": {"image/jpeg"},
    ".jpeg": {"image/jpeg"},
    ".webp": {"image/webp"},
}

SUPPORTED_TYPES = DOCUMENT_TYPES | IMAGE_TYPES


def _get_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def _is_supported(filename: str, *, allow_images: bool = True) -> bool:
    ext = _get_extension(filename)
    supported = SUPPORTED_TYPES if allow_images else DOCUMENT_TYPES
    return ext in supported


def _validate_file(filename: str, size: int | None, content_type: str | None, *, allow_images: bool = True) -> str | None:
    """Return error message if invalid, None if OK."""
    if size is not None and size > MAX_FILE_SIZE:
        return "File exceeds 100MB limit"
    if not _is_supported(filename, allow_images=allow_images):
        ext = _get_extension(filename)
        return f"Unsupported file type: {ext}"
    return None


# 文件头魔数：校验真实内容与扩展名一致，防扩展名欺骗（如 evil.php 改名 evil.png 上传）。
MAGIC_BYTES: dict[str, bytes] = {
    ".png": b"\x89PNG\r\n\x1a\n",
    ".jpg": b"\xff\xd8\xff",
    ".jpeg": b"\xff\xd8\xff",
    ".webp": b"RIFF",
    ".pdf": b"%PDF-",
    ".docx": b"PK\x03\x04",
    ".xlsx": b"PK\x03\x04",
}


def matches_magic(filename: str, content: bytes) -> bool:
    """文件头魔数是否与扩展名一致。未登记魔数的扩展名保守放行。"""
    ext = _get_extension(filename)
    if ext == ".webp":
        return content[:4] == b"RIFF" and content[8:12] == b"WEBP"
    expected = MAGIC_BYTES.get(ext)
    if expected is None:
        return True
    return content.startswith(expected)


async def save_file(file: UploadFile, *, allow_images: bool = True, user_id: int | str = 1) -> tuple[str, str]:
    """Save uploaded file (streamed to disk) and return (stored_filename, original_name)."""
    filename = file.filename or "file"
    error = _validate_file(filename, file.size, file.content_type, allow_images=allow_images)
    if error:
        raise ValueError(error)

    ext = _get_extension(filename)
    stored_name = f"{uuid.uuid4().hex}{ext}"
    user_dir = get_user_upload_dir(user_id)
    user_dir.mkdir(parents=True, exist_ok=True)
    path = user_dir / stored_name

    # 流式写盘：1MB chunk 边读边写，避免一次性 read 全文件撑爆内存；
    # 第一个 chunk 做魔数校验，累计 size 防超限，任何失败清理半成品。
    total = 0
    checked_magic = False
    output = await asyncio.to_thread(path.open, "wb")
    success = False
    try:
        while True:
            chunk = await file.read(1024 * 1024)
            if not chunk:
                break
            if not checked_magic:
                if not matches_magic(filename, chunk):
                    raise ValueError("文件内容与扩展名不符")
                checked_magic = True
            total += len(chunk)
            if total > MAX_FILE_SIZE:
                raise ValueError("File exceeds 100MB limit")
            await asyncio.to_thread(output.write, chunk)
        if total == 0:
            raise ValueError("Uploaded file is empty")
        success = True
    finally:
        await asyncio.to_thread(output.close)
        if not success:
            await asyncio.to_thread(path.unlink, missing_ok=True)
    return stored_name, filename


def delete_uploaded_file(stored_name: str, user_id: int) -> bool:
    """Delete a single uploaded file by stored name and user directory. Returns True if deleted."""
    path = get_uploaded_file_path(user_id, stored_name)
    try:
        if path.exists() and path.is_file():
            path.unlink()
            logger.info("Deleted uploaded file: %s", path)
            return True
    except Exception:
        logger.exception("Failed to delete uploaded file: %s", path)
    return False


async def is_hidden_soft_deleted_file(
    db: AsyncSession,
    *,
    filename: str,
    user_id: int,
    username: str,
) -> bool:
    """仅当文件只被软删除的文件库记录引用时隐藏。"""
    from src.database.models import BlogPost, Conversation, FileDocument, Message

    canonical_path = normalize_workspace_file_path(filename)
    legacy_name = PurePosixPath(canonical_path).name
    states = (
        await db.execute(
            select(FileDocument.deleted_at).where(
                FileDocument.user_id == str(user_id),
                FileDocument.file_path.in_({filename, canonical_path, legacy_name}),
            )
        )
    ).scalars().all()
    if not states or any(deleted_at is None for deleted_at in states):
        return False

    references = {
        filename,
        canonical_path,
        legacy_name,
        f"/api/v1/uploads/{filename}",
        f"/api/v1/uploads/{canonical_path}",
        f"/api/v1/public/uploads/{username}/{filename}",
        f"/api/v1/public/uploads/{username}/{canonical_path}",
        f"/api/v1/blog/cover/{filename}",
        f"/api/v1/blog/cover/{canonical_path}",
    }
    message_ref = await db.execute(
        select(Message.id)
        .join(Conversation, Conversation.id == Message.conversation_id)
        .where(
            Conversation.user_id == user_id,
            Conversation.deleted_at.is_(None),
            or_(Message.image_url.in_(references), Message.file_url.in_(references)),
        )
        .limit(1)
    )
    if message_ref.first() is not None:
        return False

    blog_ref = await db.execute(
        select(BlogPost.id).where(
            BlogPost.user_id == user_id,
            BlogPost.deleted_at.is_(None),
            BlogPost.cover_image.in_(references),
        ).limit(1)
    )
    return blog_ref.first() is None


async def vectorize_and_store(
    stored_filename: str,
    collection_name: str,
    original_name: str | None = None,
    user_id: int | str = 1,
    resource_type: str | None = None,
    resource_id: int | None = None,
    progress_reporter=None,
) -> list[str]:
    """Parse, chunk, embed, and store documents; returns the stored chunk contents."""
    from src.services.infra.embeddings.embedding_service import get_embeddings

    from src.utils.chunker import chunk_text

    async def report(stage: str, completed: int, total: int, unit: str) -> None:
        if progress_reporter:
            await progress_reporter(stage, completed, total, unit)

    _t0 = time.time()
    logger.info(
        "KB vectorization parse started: stored_name=%s original_name=%s user_id=%s",
        stored_filename,
        original_name,
        user_id,
    )
    loop = asyncio.get_running_loop()

    def parser_progress(completed: int, total: int, unit: str) -> None:
        if progress_reporter:
            asyncio.run_coroutine_threadsafe(report("parse", completed, total, unit), loop).result()

    path = get_uploaded_file_path(int(user_id), stored_filename)
    await report("parse", 0, 1, "operation")
    text = await asyncio.to_thread(file_parser.parse_path, path, parser_progress)
    logger.info(
        "KB vectorization parse completed: stored_name=%s text_length=%s",
        stored_filename,
        len(text),
    )

    await report("chunk", 0, 1, "operation")
    chunks = await asyncio.to_thread(chunk_text, text)
    await report("chunk", 1, 1, "operation")
    logger.info(
        "KB vectorization chunking completed: stored_name=%s chunks=%s",
        stored_filename,
        len(chunks),
    )
    if not chunks:
        logger.warning("KB vectorization produced no chunks: stored_name=%s", stored_filename)
        raise ValueError("Document contains no indexable text")

    logger.info("KB vectorization embedding started: stored_name=%s chunks=%s", stored_filename, len(chunks))

    async def embedding_progress(completed: int, total: int, unit: str) -> None:
        await report("embedding", completed, total, unit)

    await report("embedding", 0, len(chunks), "chunk")
    embeddings = await get_embeddings(chunks, progress_callback=embedding_progress)
    if len(embeddings) != len(chunks):
        raise ValueError(f"Embedding count mismatch: chunks={len(chunks)} embeddings={len(embeddings)}")
    logger.info("KB vectorization embedding completed: stored_name=%s embeddings=%s", stored_filename, len(embeddings))

    file_name = original_name or stored_filename
    file_type = _get_extension(file_name).lstrip(".")
    metadata_list = []
    await report("metadata", 0, len(chunks), "chunk")
    for index in range(len(chunks)):
        metadata = {
            "source": file_name,
            "file_name": file_name,
            "original_name": file_name,
            "stored_name": stored_filename,
            "file_type": file_type,
            "collection_name": collection_name,
            "user_id": user_id,
            "resource_type": resource_type,
            "resource_id": resource_id,
            "chunk_index": index,
            "total_chunks": len(chunks),
            "chunk_id": f"{stored_filename}:{index}",
        }
        metadata_list.append(metadata)
        await report("metadata", index + 1, len(chunks), "chunk")

    logger.info(
        "KB vectorization store started: stored_name=%s collection=%s chunks=%s",
        stored_filename,
        collection_name,
        len(chunks),
    )
    async def vector_progress(completed: int, total: int, unit: str) -> None:
        await report("vector_store", completed, total, unit)

    await report("vector_store", 0, len(chunks), "chunk")
    await _store_document_chunks(
        collection_name=collection_name,
        chunks=chunks,
        metadata_list=metadata_list,
        embeddings=embeddings,
        user_id=user_id,
        progress_callback=vector_progress,
    )
    logger.info(
        "KB vectorization store completed: stored_name=%s collection=%s chunks=%s",
        stored_filename,
        collection_name,
        len(chunks),
    )
    logger.info(
        "KB vectorization done: stored_name=%s user_id=%s chunks=%d duration_ms=%d",
        stored_filename, user_id, len(chunks), int((time.time() - _t0) * 1000),
    )

    return chunks


async def vectorize_text_and_store(
    text: str,
    collection_name: str,
    *,
    source_id: str,
    original_name: str | None = None,
    user_id: int | str = 1,
    resource_type: str = "blog_post",
    resource_id: int | None = None,
    progress_reporter=None,
) -> list[str]:
    """对纯文本（如博客 Markdown 正文）分块、向量化、写入向量库；跳过文件解析，metadata 的 stored_name/chunk_id 用 source_id（资源稳定标识，如 "blog_post:123"），附带 resource_type 便于检索过滤。"""
    from src.services.infra.embeddings.embedding_service import get_embeddings
    from src.utils.chunker import chunk_text

    async def report(stage: str, completed: int, total: int, unit: str) -> None:
        if progress_reporter:
            await progress_reporter(stage, completed, total, unit)

    logger.info(
        "KB text vectorization started: source_id=%s resource_type=%s user_id=%s",
        source_id,
        resource_type,
        user_id,
    )

    # 对齐 index_v1 权重表（含 parse）：文本无需解析，瞬间完成 parse 段。
    await report("parse", 1, 1, "operation")

    await report("chunk", 0, 1, "operation")
    chunks = await asyncio.to_thread(chunk_text, text)
    await report("chunk", 1, 1, "operation")
    logger.info(
        "KB text vectorization chunking completed: source_id=%s chunks=%s",
        source_id,
        len(chunks),
    )
    if not chunks:
        logger.warning("KB text vectorization produced no chunks: source_id=%s", source_id)
        raise ValueError("Document contains no indexable text")

    async def embedding_progress(completed: int, total: int, unit: str) -> None:
        await report("embedding", completed, total, unit)

    await report("embedding", 0, len(chunks), "chunk")
    embeddings = await get_embeddings(chunks, progress_callback=embedding_progress)
    if len(embeddings) != len(chunks):
        raise ValueError(f"Embedding count mismatch: chunks={len(chunks)} embeddings={len(embeddings)}")

    display_name = original_name or source_id
    metadata_list = []
    await report("metadata", 0, len(chunks), "chunk")
    for index in range(len(chunks)):
        metadata = {
            "source": display_name,
            "original_name": display_name,
            "stored_name": source_id,
            "resource_type": resource_type,
            "resource_id": resource_id,
            "collection_name": collection_name,
            "user_id": user_id,
            "chunk_index": index,
            "total_chunks": len(chunks),
            "chunk_id": f"{source_id}:{index}",
        }
        metadata_list.append(metadata)
        await report("metadata", index + 1, len(chunks), "chunk")

    async def vector_progress(completed: int, total: int, unit: str) -> None:
        await report("vector_store", completed, total, unit)

    await report("vector_store", 0, len(chunks), "chunk")
    await _store_document_chunks(
        collection_name=collection_name,
        chunks=chunks,
        metadata_list=metadata_list,
        embeddings=embeddings,
        user_id=user_id,
        progress_callback=vector_progress,
    )
    logger.info(
        "KB text vectorization completed: source_id=%s collection=%s chunks=%s",
        source_id,
        collection_name,
        len(chunks),
    )
    return chunks


# ---- File Preview (docx/xlsx → HTML) ----


async def convert_to_html(stored_filename: str, *, user_id: int | str = 1) -> str:
    """Convert a stored file to HTML for browser preview; PDF is served as-is (browser renders natively)."""
    path = get_uploaded_file_path(int(user_id), stored_filename)
    ext = _get_extension(stored_filename)
    if ext == ".docx":
        return _docx_to_html(path)
    elif ext == ".xlsx":
        return _xlsx_to_html(path)
    else:
        raise ValueError(f"Preview not supported for: {ext}")


def _docx_to_html(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    html_parts: list[str] = ['<div class="docx-preview">']
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            html_parts.append("<p>&nbsp;</p>")
            continue
        style = para.style.name.lower() if para.style else ""
        if "heading 1" in style:
            html_parts.append(f"<h1>{_escape_html(text)}</h1>")
        elif "heading 2" in style:
            html_parts.append(f"<h2>{_escape_html(text)}</h2>")
        elif "heading 3" in style:
            html_parts.append(f"<h3>{_escape_html(text)}</h3>")
        else:
            html_parts.append(f"<p>{_escape_html(text)}</p>")
    for table in doc.tables:
        html_parts.append('<table border="1" cellpadding="4" style="border-collapse:collapse;width:100%">')
        for i, row in enumerate(table.rows):
            tag = "th" if i == 0 else "td"
            html_parts.append(f"<tr>{''.join(f'<{tag}>{_escape_html(cell.text.strip())}</{tag}>' for cell in row.cells)}</tr>")
        html_parts.append("</table>")
    html_parts.append("</div>")
    return "\n".join(html_parts)


def _xlsx_to_html(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    html_parts: list[str] = ['<div class="xlsx-preview">']
    for ws in wb.worksheets:
        html_parts.append(f"<h3>{_escape_html(ws.title)}</h3>")
        html_parts.append('<table border="1" cellpadding="4" style="border-collapse:collapse;width:100%">')
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            tag = "th" if row_idx == 0 else "td"
            cells = "".join(f"<{tag}>{_escape_html(str(c) if c is not None else '')}</{tag}>" for c in row)
            html_parts.append(f"<tr>{cells}</tr>")
        html_parts.append("</table>")
    html_parts.append("</div>")
    wb.close()
    return "\n".join(html_parts)


def _escape_html(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
