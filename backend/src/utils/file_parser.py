import re
from pathlib import Path


def _get_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def parse_path(path: Path, progress_callback=None) -> str:
    """Synchronously parse a path, optionally reporting completed parser units."""
    ext = _get_extension(path.name)
    if ext == ".docx":
        return _parse_docx(path, progress_callback)
    if ext == ".xlsx":
        return _parse_xlsx(path, progress_callback)
    if ext == ".pdf":
        return _parse_pdf(path, progress_callback)
    raise ValueError(f"Cannot parse file type: {ext}")


def _parse_docx(path: Path, progress_callback=None) -> str:
    from docx import Document

    try:
        doc = Document(str(path))
    except ValueError as e:
        raise ValueError(f"Invalid or corrupted DOCX file: {e}") from e

    total = len(doc.paragraphs) + sum(len(table.rows) for table in doc.tables) + 1
    completed = 0
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            style_name = para.style.name.lower() if para.style else ""
            if "heading 1" in style_name:
                parts.append(f"\n\n## {text}")
            elif "heading 2" in style_name:
                parts.append(f"\n\n### {text}")
            elif "heading 3" in style_name:
                parts.append(f"\n\n#### {text}")
            else:
                parts.append(text)
        completed += 1
        if progress_callback:
            progress_callback(completed, total, "block")

    for table in doc.tables:
        parts.append("")
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
            completed += 1
            if progress_callback:
                progress_callback(completed, total, "block")
        parts.append("")

    if progress_callback:
        progress_callback(total, total, "block")
    return "\n".join(parts).strip()


def _parse_xlsx(path: Path, progress_callback=None) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    total = len(wb.worksheets) + 1
    parts: list[str] = []
    for index, ws in enumerate(wb.worksheets, start=1):
        parts.append(f"\n\n=== Sheet: {ws.title} ===")
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            cells = [str(c) if c is not None else "" for c in row]
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                parts.append(f"### {row_text}" if row_idx == 0 else row_text)
        parts.append("")
        if progress_callback:
            progress_callback(index, total, "sheet")
    wb.close()
    if progress_callback:
        progress_callback(total, total, "sheet")
    return "\n".join(parts).strip()


def _parse_pdf(path: Path, progress_callback=None) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    total = len(reader.pages) + 1
    parts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text()
        if text:
            parts.append(f"[--- 第{index}页 ---]\n{text}")
        if progress_callback:
            progress_callback(index, total, "page")
    result = _clean_pdf_text("\n".join(parts))
    if progress_callback:
        progress_callback(total, total, "page")
    return result


def _clean_pdf_text(text: str) -> str:
    lines = text.split("\n")
    cleaned: list[str] = []

    for line in lines:
        stripped = line.rstrip()
        if stripped.startswith("[--- 第") and stripped.endswith("页 ---]"):
            cleaned.append(stripped)
            continue
        if re.fullmatch(r"第\s*\d+\s*页|Page\s+\d+|-\s*\d+\s*-|\s*\d+\s*", stripped, re.IGNORECASE):
            continue
        if cleaned and stripped and stripped == cleaned[-1].strip():
            continue
        cleaned.append(stripped)

    return re.sub(r"\n{3,}", "\n\n", "\n".join(cleaned)).strip()
